# -*- coding: utf-8 -*-
"""一次性回填：为三张台账的「项目编号」列补历史值（幂等，只填空格）。
项目编号解析：归档路径反解 > 关键词/合同名/事由 找既有项目夹。
用法：python tools/backfill_projcode.py [--dry-run]
"""
import argparse
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))
from oa_common import paths, moneyline  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

JOBS = [
    # (台账路径, sheet, 项目编号列, 关键词列, 路径列)
    (paths.LEDGER_CONTRACT, "合同台账", 15, 7, 10),  # 关联事项审批关键词 / 合同文本路径
    (paths.LEDGER_PAY, "付款台账", 15, 2, 10),       # 合同名称 / 发票路径
    (paths.LEDGER_OA, "发起台账", 17, 2, None),      # 事由
]


def strip_contract(cn):
    for suf in ("采购合同", "服务合同", "合同"):
        if cn.endswith(suf) and len(cn) > len(suf):
            return cn[:-len(suf)]
    return cn


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    for p, sheet, col, kcol, pcol in JOBS:
        wb = load_workbook(p)
        ws = wb[sheet]
        n = 0
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, col).value or "").strip():
                continue
            kw = str(ws.cell(r, kcol).value or "").strip()
            if p.name.startswith("付款"):
                kw = strip_contract(kw)
            path = str(ws.cell(r, pcol).value or "") if pcol else ""
            code, matter = moneyline.resolve_code(kw, path)
            if code and "示例" not in str(ws.cell(r, 1).value or "") + kw:
                print(f"[{p.name}] row{r} {kw or path[:30]} -> {code}")
                if not a.dry_run:
                    ws.cell(r, col, code)
                n += 1
        if not a.dry_run:
            wb.save(p)
        wb.close()
        print(f"== {p.name}: {'回填' if n else '无需回填'} {n} 行"
              + ("（dry-run 未保存）" if a.dry_run else ""))


if __name__ == "__main__":
    main()
