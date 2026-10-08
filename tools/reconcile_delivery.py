# -*- coding: utf-8 -*-
"""交付台账对账：交付区实物 PDF ↔ 台账登记行 账实相符。

扫描 交付/<功能>/*.pdf（排除台账自身），与 记录/台账/交付台账.xlsx（主份）及
交付/交付台账.xlsx（副本）逐行比对。
  · 默认（追加模式）：任一分身缺行的册子补记一行（成员清单从
    运行/state/delivery_manifest.json 反查，查不到留空备注说明）。幂等。
  · --rebuild：以交付区实物为唯一基准整表重建两份分身——一本一行，
    彻底清除改名/重跑留下的幽灵行；生成时间沿用旧台账里该册最早的记录，
    旧账没有的取册子文件修改时间。重建前自动把旧表备份到 运行/state/。

用法：python -m tools.reconcile_delivery [--dry-run] [--rebuild]
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from oa_common import paths, deliver  # noqa: E402

LEDGER_HEAD = deliver.LEDGER_HEAD


def _existing_rows(led):
    """返回 {(功能, 交付文件): 行号}；台账不存在返回空。"""
    from openpyxl import load_workbook
    if not Path(led).exists():
        return {}
    wb = load_workbook(led, read_only=True)
    ws = wb.active
    idx = {}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or len(row) < 4:
            continue
        idx[(str(row[1] or ""), str(row[3] or ""))] = r
    wb.close()
    return idx


def _members_from_manifest(func, matter):
    try:
        man = json.loads(deliver.MANIFEST.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return ""
    for key, members in man.items():
        if key.startswith("_") or "|" not in key:
            continue
        f, m = key.split("|", 1)
        if f == func and m == matter:
            return "、".join(Path(p).name for p in sorted(members, key=lambda x: Path(x).name))
    return ""


def _booklets():
    for d in sorted(paths.DELIVER.iterdir()):
        if not d.is_dir():
            continue
        for p in sorted(d.glob("*.pdf")):
            yield d.name, p


def _old_times():
    """两份旧台账里每个（功能,文件）最早出现的生成时间，重建时沿用。"""
    from openpyxl import load_workbook
    best = {}
    for led in deliver.LEDGERS:
        if not Path(led).exists():
            continue
        wb = load_workbook(led, read_only=True)
        for row in wb.active.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 4:
                continue
            k = (str(row[1] or ""), str(row[3] or ""))
            t = str(row[0] or "")
            if k not in best or (t and t < best[k]):
                best[k] = t
        wb.close()
    return best


def _matter_of(name):
    import re
    m = re.search(r"（(.+?)(?:·金额|$)", name)
    return m.group(1) if m else Path(name).stem


def _rebuild(dry):
    import shutil
    from openpyxl import Workbook
    times = _old_times()
    rows = []
    for func, p in _booklets():
        import datetime as _dt
        t = times.get((func, p.name)) or _dt.datetime.fromtimestamp(
            p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
        members = _members_from_manifest(func, _matter_of(p.name))
        mem = [x for x in members.split("、") if x]
        pages = ""
        try:
            import pymupdf as fitz
            with fitz.open(str(p)) as doc:
                pages = doc.page_count
        except Exception:
            pass
        rows.append([t, func, _matter_of(p.name), p.name, len(mem), pages, members, "重建"])
    rows.sort(key=lambda r: (r[1], r[3]))
    print(f"重建 {len(rows)} 行（旧主台账幽灵行一并清除）")
    if dry:
        for r in rows:
            print("  ", r[1], r[3])
        return
    bak = paths.STATE / f"交付台账备份_{dt.datetime.now():%Y%m%d_%H%M%S}"
    bak.mkdir(parents=True, exist_ok=True)
    for led in deliver.LEDGERS:
        if Path(led).exists():
            shutil.copy2(led, bak / f"{Path(led).parent.name}_{Path(led).name}")
    for led in deliver.LEDGERS:
        wb = Workbook()
        ws = wb.active
        ws.title = "交付台账"
        ws.append(LEDGER_HEAD)
        for r in rows:
            ws.append(r)
        Path(led).parent.mkdir(parents=True, exist_ok=True)
        wb.save(led)
        print(f"已重写 {led}（{len(rows)}行），旧表备份在 {bak}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--rebuild", action="store_true", help="以交付区实物为基准整表重建两份分身")
    a = ap.parse_args()
    if a.rebuild:
        _rebuild(a.dry_run)
        return

    from openpyxl import Workbook, load_workbook
    have = {str(led): _existing_rows(led) for led in deliver.LEDGERS}
    n = 0
    for func, p in _booklets():
        for led in deliver.LEDGERS:
            if (func, p.name) in have[str(led)]:
                continue
            m = __import__("re").search(r"（(.+?)(?:·金额|$)", p.name)
            matter = m.group(1) if m else p.stem
            row = [dt.datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M"),
                   func, matter, p.name, "", "",
                   _members_from_manifest(func, matter), "对账补记"]
            try:
                import pymupdf as fitz
                with fitz.open(str(p)) as doc:
                    row[4], row[5] = len([x for x in row[6].split("、") if x]), doc.page_count
            except Exception:
                pass
            print(f"[补记] {Path(led).name} <- {func}/{p.name}")
            if not a.dry_run:
                if Path(led).exists():
                    wb = load_workbook(led)
                    ws = wb.active
                else:
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "交付台账"
                    ws.append(LEDGER_HEAD)
                ws.append(row)
                Path(led).parent.mkdir(parents=True, exist_ok=True)
                wb.save(led)
                have[str(led)][(func, p.name)] = -1
            n += 1
    print(f"{'（试运行）' if a.dry_run else ''}共需补记 {n} 行")


if __name__ == "__main__":
    main()
