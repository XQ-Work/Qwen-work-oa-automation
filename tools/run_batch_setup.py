# -*- coding: utf-8 -*-
"""批量准备：对指定项目 自动发现合同->extract --commit->按场景补B列/拟签。
打印 项目->台账行号，供随后 fill --save 使用。不跑浏览器。"""
import sys
import subprocess
import re
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))
sys.path.insert(0, str(BASE / "modules" / "m05_contract"))
from oa_common import paths  # noqa: E402
import extract_contract as ex  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ARCH = paths.DATA / "OA附件库/文档整理/采购项目档案"
SCENE_MAP = {"货物": "货物-生产物资/设备", "服务": "服务-维修/外包"}
EXE = BASE / "modules/m05_contract/extract_contract.py"
PROJECTS = ["2026-003-潜污泵采购", "2026-005-发电机排涝采购",
            "2026-007-细格栅维修", "2026-009-垃圾清运服务"]


def find_contract(proj):
    d = ARCH / proj / "04_采购合同"
    cs = [f for f in d.glob("*.pdf")
          if not any(s in f.name for s in ("已用印", "盖章", "扫描"))]
    return cs[0] if cs else None


for proj in PROJECTS:
    p = find_contract(proj)
    if not p:
        print(f"[{proj}] 无合同，跳过"); continue
    # 场景
    scene = ex.detect_scene(ex.read_text(p), p.stem)
    # extract --commit
    r = subprocess.run([sys.executable, str(EXE), str(p), "--commit"],
                       capture_output=True, text=True, encoding="utf-8")
    m = re.search(r"(新增|更新)台账第(\d+)行", r.stdout or "")
    if not m:
        print(f"[{proj}] commit 失败:\n", (r.stdout or '')[-300:] + (r.stderr or '')[-300:])
        continue
    row = int(m.group(2))
    wb = load_workbook(paths.DATA / "合同发起台账.xlsx")
    ws = wb["合同台账"]
    ws.cell(row, 2, SCENE_MAP[scene]); ws.cell(row, 8, "2026-09-18")
    wb.save(paths.DATA / "合同发起台账.xlsx")
    print(f"[{proj}] -> 行{row} 场景={SCENE_MAP[scene]} "
          f"乙方={ws.cell(row,6).value} 金额={ws.cell(row,4).value}")
