# -*- coding: utf-8 -*-
"""m06 付款台账录入：解析电子发票 -> 追加一行到《付款发起台账.xlsx》。

用法：
  python intake_payment.py <发票.pdf> --contract 空压机采购合同 \
      --item-type 货物类采购事项 [--commit]
"""
import argparse
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE))
sys.path.insert(0, str(BASE / "modules" / "m06_payment"))
from oa_common import paths  # noqa: E402
import parse_invoice as pi  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LEDGER = paths.DATA / "付款发起台账.xlsx"
SHEET = "付款台账"
# 列: 2合同名称 3付款单位 4收款单位 5付款金额 6采购类事项类型
#     7发票号码 8开票日期 9费用说明 10发票路径 11状态


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("invoice")
    ap.add_argument("--contract", default="", help="合同名称")
    ap.add_argument("--item-type", default="货物类采购事项")
    ap.add_argument("--payer", default="荆州浦华荆清水务有限公司")
    ap.add_argument("--commit", action="store_true", help="追加一行到台账")
    args = ap.parse_args()

    p = Path(args.invoice)
    if not p.exists():
        raise SystemExit(f"发票不存在：{p}")
    inv = pi.parse(p)
    amt = inv["价税合计"]
    desc = f"{args.contract}，已完成入库验收，按照合同约定付款{amt}元"
    print("=" * 55)
    print(f"发票：{p.name}")
    print(f"  收款单位={inv['销售方']}  付款金额={amt}  发票号码={inv['发票号码']}  开票日期={inv['开票日期']}")
    print(f"  费用说明={desc}")
    if not inv["销售方"] or not amt:
        print("[!] 收款方或金额未解析出，请核对发票（可能扫描件，需视觉兜底）")
        return
    if not args.commit:
        print("（预览，未写入。加 --commit 追加台账行）")
        return
    wb = load_workbook(LEDGER)
    ws = wb[SHEET]
    r = ws.max_row + 1
    vals = {2: args.contract, 3: args.payer, 4: inv["销售方"], 5: amt,
            6: args.item_type, 7: inv["发票号码"], 8: inv["开票日期"],
            9: desc, 10: str(p.resolve()), 11: "待确认"}
    for c, v in vals.items():
        if v:
            ws.cell(r, c, v)
    wb.save(LEDGER)
    print(f"[ok] 已追加台账第{r}行（状态=待确认，核对后跑 fill_payment --row {r}）")


if __name__ == "__main__":
    main()
