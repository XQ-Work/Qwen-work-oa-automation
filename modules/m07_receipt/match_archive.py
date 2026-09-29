# -*- coding: utf-8 -*-
"""M07 转账回单 → 匹配台账 → 归档到项目档案 → 回写。

流程：读《付款回单台账》回单记录（收款人/金额/摘要/流水号）
  → 找回单图片(已完成\\{收款人}\\…_{金额}_…)
  → 依次匹配：付款发起台账(收款单位+金额) → 合同发起台账(乙方+金额)
  → 命中则把回单图复制到 采购项目档案\\{项目编号}-{事项}\\05_付款及验收\\
  → 命中付款台账行则回写 状态=已付款 + 备注(回单文件)
匹配不上的留待人工。幂等：已处理过的会计流水号记录在 项目归档记录.json。

用法：
  python match_archive.py            # 匹配+归档+回写
  python match_archive.py --dry-run  # 只预览匹配结果，不复制不回写
"""
import argparse
import json
import re
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE))
sys.path.insert(0, str(BASE / "modules" / "m06_payment"))
from oa_common import paths  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

回单目录 = paths.RECEIPT_DIR
回单台账 = 回单目录 / "付款回单台账.xlsx"
档案根 = paths.ARCHIVE
LOG = 回单目录 / "项目归档记录.json"


def line0(rec):
    return f"[{rec['日期']}] {rec['收款人']} {rec['金额']}元 «{rec['摘要']}»"


def _f(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def _proj(合同名称):
    """合同名称 → 项目名（去尾部 采购合同/服务合同/合同）。"""
    cn = str(合同名称 or "").strip()
    for suf in ("采购合同", "服务合同", "合同"):
        if cn.endswith(suf) and len(cn) > len(suf):
            return cn[:-len(suf)]
    return cn


def load_receipts():
    if not 回单台账.exists():
        print("回单台账不存在（尚未扫描过银行回单？）:", 回单台账)
        return []
    wb = load_workbook(回单台账, data_only=True)
    ws = wb["付款回单台账"]
    out = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 3).value
        amt = _f(ws.cell(r, 5).value)
        if not name or amt is None:
            continue
        out.append({"row": r, "日期": str(ws.cell(r, 2).value or ""),
                    "收款人": str(name).strip(), "金额": f"{amt:.2f}",
                    "摘要": str(ws.cell(r, 7).value or ""),
                    "流水号": str(ws.cell(r, 8).value or "")})
    wb.close()
    return out


def receipt_image(收款人, 金额):
    d = 回单目录 / "已完成" / 收款人
    if not d.exists():
        return None
    for f in sorted(list(d.glob("*.png")) + list(d.glob("*.jpg"))):
        if f"_{金额}_" in f.name:
            return f
    return None


def find_project_dir(项目名):
    if not 项目名:
        return None
    for d in sorted(档案根.iterdir()):
        if d.is_dir() and 项目名 in d.name:
            return d
    return None


def _is_sample(ws, r):
    """示例行跳过（序号/首列/场景名 含“示例”）。"""
    for c in (1, 2):
        if "示例" in str(ws.cell(r, c).value or ""):
            return True
    return False


def lane_b_folders(摘要):
    """Lane B(<2000 无合同)：回单摘要里含哪个项目文件夹的事项名，就归哪个项目夹。
    要求转账时把事项名写进摘要（今后约定）。返回 [(文件夹, 事项名)]。"""
    if not 摘要:
        return []
    hits = []
    for d in sorted(档案根.iterdir()):
        if not d.is_dir():
            continue
        m = re.match(r"\d{4}-\d{3}-(.+)", d.name)
        if m and len(m.group(1)) >= 2 and m.group(1) in 摘要:
            hits.append((d, m.group(1)))
    return hits


def match_ledgers(name, amt):
    """金额链容差匹配：完全相等 > ±1元尾差(自动+备注) > ±5%(仅提示待人工)。
    候选先付款台账后合同台账，同级优先能解析到项目文件夹的行。
    返回 (来源, 行号, 项目名, 备注, ±5%候选列表[(来源,行号,项目名,差额)])。"""
    amt_f = _f(amt)
    if amt_f is None:
        return (None, 0, "", [], [])
    cands = []           # (差额, 优先级0付款/1合同, 来源, 行号, 项目名)
    wb = load_workbook(paths.LEDGER_PAY, data_only=True)
    ws = wb["付款台账"]
    for r in range(2, ws.max_row + 1):
        if _is_sample(ws, r) or str(ws.cell(r, 4).value or "").strip() != name:
            continue
        ra = _f(ws.cell(r, 5).value)
        if ra is None:
            continue
        if "已付款" in str(ws.cell(r, 11).value or ""):
            continue                         # 该笔已有回单
        cands.append((round(abs(ra - amt_f), 2), 0, "付款台账", r, _proj(ws.cell(r, 2).value)))
    wb.close()
    if not any(c[1] == 0 and c[0] <= 1.0 for c in cands):
        wb = load_workbook(paths.LEDGER_CONTRACT, data_only=True)
        ws = wb["合同台账"]
        for r in range(2, ws.max_row + 1):
            if _is_sample(ws, r) or str(ws.cell(r, 6).value or "").strip() != name:
                continue
            ra = _f(ws.cell(r, 4).value)
            if ra is None:
                continue
            cands.append((round(abs(ra - amt_f), 2), 1, "合同台账", r, _proj(ws.cell(r, 3).value)))
        wb.close()
    near5 = [(src, row, proj, d) for d, _, src, row, proj in cands if 1.0 < d <= max(amt_f * 0.05, 1.01)]
    auto = [c for c in cands if c[0] <= 1.0]
    if not auto:
        return (None, 0, "", [], near5)
    exact = [c for c in auto if c[0] <= 0.005]
    pool = exact or auto
    pool.sort(key=lambda c: (c[0], c[1]))
    best = next((c for c in pool if find_project_dir(c[4])), pool[0])
    note = "" if best[0] <= 0.005 else f"尾差{best[0]}元"
    return (best[2], best[3], best[4], note, near5)


def writeback_pay(row, img_name):
    wb = load_workbook(paths.LEDGER_PAY)
    ws = wb["付款台账"]
    ws.cell(row, 11, "已付款")
    note = str(ws.cell(row, 14).value or "").strip()
    ws.cell(row, 14, (note + "；" if note else "") + f"回单:{img_name}")
    wb.save(paths.LEDGER_PAY)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="只预览匹配，不复制不回写")
    args = ap.parse_args()

    log = json.loads(LOG.read_text(encoding="utf-8")) if LOG.exists() else {}
    # 幂等键：会计流水号为空时用内容指纹，避免多条空流水号互相挤掉
    def key_of(rec):
        return rec["流水号"] or f"{rec['收款人']}|{rec['金额']}|{rec['摘要']}|{rec['日期']}"
    receipts = load_receipts()
    print(f"回单记录 {len(receipts)} 条，已归档 {len(log)} 条")
    n_ok = n_manual = 0
    for rec in receipts:
        k = key_of(rec)
        if k in log:
            continue
        img = receipt_image(rec["收款人"], rec["金额"])
        src, row, proj, note, near5 = match_ledgers(rec["收款人"], rec["金额"])
        pdir = None
        line = line0(rec)
        # Lane A：合同/付款台账命中
        if src:
            line += f" → {src}#{row} 项目名={proj}" + (f" ({note})" if note else "")
            pdir = find_project_dir(proj)
        else:
            # Lane B：<2000 无合同，按摘要里的事项名归项目夹（转账时须写明事项名）
            lb = lane_b_folders(rec["摘要"])
            if len(lb) == 1:
                pdir, proj = lb[0]
                src, row = "事项审批(摘要)", 0
                line = line0(rec) + f" → LaneB 项目名={proj}"
            elif len(lb) > 1:
                n_manual += 1
                print(line0(rec) + f"  [Lane B歧义]{[p.name for p, _ in lb]}")
                continue
            elif near5:
                n_manual += 1
                tips = "；".join(f"{s}#{r} {p} 差{d}元" for s, r, p, d in near5[:4])
                print(line0(rec) + f"  [±5%尾差待人工] {tips}")
                continue
        if not src or not img:
            n_manual += 1
            print(line + ("" if img else "  [缺回单图]"))
            continue
        if not pdir:
            n_manual += 1
            print(line + "  [未找到项目文件夹]")
            continue
        dest = pdir / "05_付款及验收" / img.name
        print(line + f"  → 归档 {dest.relative_to(paths.DATA)}")
        if not args.dry_run:
            dest.parent.mkdir(parents=True, exist_ok=True)
            if not dest.exists():
                shutil_copy(img, dest)
            if src == "付款台账":
                writeback_pay(row, img.name)
            try:    # 金额链主档：累计已付+差额重算（Lane B 也会补记已付款行）
                from oa_common import moneyline
                code, matter = moneyline.code_from_path(pdir.name)
                if code:
                    moneyline.mark_paid(code, rec["金额"], rec["收款人"],
                                        receipt_note=img.name, receipt_src=(rec["摘要"],))
            except Exception as e:
                print("  [moneyline] 主档更新失败(不影响归档): " + str(e)[:80])
            log[k] = {"归档": str(dest.relative_to(paths.DATA)),
                      "匹配": f"{src}#row{row}" + (f"({note})" if note else ""),
                      "日期": rec["日期"]}
            LOG.write_text(json.dumps(log, ensure_ascii=False, indent=1), encoding="utf-8")
        n_ok += 1
    print(f"完成：匹配归档 {n_ok} 条，待人工 {n_manual} 条"
          + ("（dry-run 未写盘）" if args.dry_run else ""))


def shutil_copy(src, dest):
    import shutil
    shutil.copy2(src, dest)


if __name__ == "__main__":
    main()
