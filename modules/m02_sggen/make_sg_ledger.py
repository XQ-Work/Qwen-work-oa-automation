# -*- coding: utf-8 -*-
"""生成《申购单需求台账.xlsx》：用户按行填物资需求 -> make_sg --ledger 读取生成申购单"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

import sys as _sys2
_sys2.path.insert(0, r"D:/自动备份/Qwen work/自动化（OA流程）/oa-automation")
from oa_common import paths as _paths
OUT = str(_paths.LEDGER_SG)

HEADERS = [
    ("序号", 6), ("事项名*", 18), ("类别*", 10), ("预估金额（元）*", 14),
    ("申购说明*", 30), ("预计使用日期", 12), ("物资名称*", 16),
    ("规格型号及技术要求", 26), ("单位", 8), ("数量", 8), ("备注", 14),
    ("状态", 10), ("生成时间", 16), ("备注2", 14),
]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
GRAY = PatternFill("solid", fgColor="D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "需求台账"
for col, (name, width) in enumerate(HEADERS, 1):
    c = ws.cell(row=1, column=col, value=name)
    c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws.column_dimensions[get_column_letter(col)].width = width
ws.row_dimensions[1].height = 30

fills = {2: YELLOW, 3: YELLOW, 4: YELLOW, 5: YELLOW, 7: YELLOW,
         6: GREEN, 8: GREEN, 9: GREEN, 10: GREEN, 11: GREEN, 14: GREEN,
         12: GRAY, 13: GRAY}

# 示例组：同一事项名两行 = 一张申购单
examples = [
    ["示例", "示例阀门申购", "货物", 1234.5, "因维修需要申购阀门一批", "2026.10",
     "DN100闸阀", "PN10", "个", "4", "", "示例", "", ""],
    ["示例", "示例阀门申购", "", "", "", "",
     "止回阀", "H44H-16 DN80", "个", "2", "备用", "示例", "", ""],
]
for r, row in enumerate(examples, 2):
    for col, val in enumerate(row, 1):
        c = ws.cell(row=r, column=col, value=val)
        c.font = Font(name="微软雅黑", size=10, color="808080", italic=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORDER

for r in range(4, 304):
    for col in range(1, len(HEADERS) + 1):
        c = ws.cell(row=r, column=col)
        c.border = BORDER
        c.font = Font(name="微软雅黑", size=10)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if col in fills:
            c.fill = fills[col]
    ws.cell(row=r, column=4).number_format = "#,##0.00"

dv = DataValidation(type="list", formula1='"货物,服务"', allow_blank=True,
                    showErrorMessage=True, error="类别只能填 货物 或 服务")
ws.add_data_validation(dv)
dv.add("C4:C303")
ws.freeze_panes = "B2"
ws.auto_filter.ref = "A1:N303"

gd = wb.create_sheet("填写说明")
gd.column_dimensions["A"].width = 100
lines = [
    ("《申购单需求台账》填写说明", True),
    ("", False),
    ("■ 一行 = 一条物资；同一申购单的多条物资 = 多行「事项名」相同。", True),
    ("■ 事项级信息（类别/预估金额/申购说明/预计使用日期）只填该事项首行，续行留空。", True),
    ("■ 黄色必填：事项名、类别（首行）、预估金额（首行）、申购说明（首行）、物资名称。", False),
    ("■ 绿色选填：规格/单位/数量尽量填全（申购单表格内容）；预估金额=整单总数（进文件名，供事项审批取数）。", False),
    ("■ 灰色两列脚本回写，不要动。示例两行（状态=示例）脚本自动跳过。", False),
    ("", False),
    ("使用方法：填好后对智能体说「生成申购单」——脚本读取所有状态为空的行，按事项名分组", True),
    ("逐单生成《申购单（事项名-金额：X）.xlsx》落进发起收件箱，并回写状态=已生成。", True),
    ("随后可直接说「跑OA」发起事项审批（②→①已自动串联）。", True),
]
for i, (text, bold) in enumerate(lines, 1):
    c = gd.cell(row=i, column=1, value=text)
    c.font = Font(name="微软雅黑", size=11 if bold else 10, bold=bold)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print("saved:", OUT)
