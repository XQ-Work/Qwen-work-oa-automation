# -*- coding: utf-8 -*-
"""申购单生成器（队列②）
输入 JSON -> 填充对应模板 -> 输出 `申购单（事项名-金额：X）.xlsx` 到发起收件箱

用法：python make_sg.py --json req.json
req.json 示例：
{
 "类别": "货物",                // 货物 | 服务
 "事项名": "PE配件",
 "金额": 999.91,                // 预估金额，写入文件名供 OA 事项审批取用
 "申购说明": "PE管及配件",
 "预计使用日期": "2026.10",     // 可选
 "明细": [
   {"名称":"PE管","规格":"200","单位":"米","数量":"12","备注":""},
   {"名称":"PE弯头","规格":"200","单位":"个","数量":"5","备注":""}
 ]
}
"""
import argparse
import datetime
import json
import re
import sys
from pathlib import Path
import sys as _sys
_sys.path.insert(0, r"D:/自动备份/Qwen work/自动化（OA流程）/系统/oa-automation")
from oa_common import paths

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string

BASE = paths.STATE
TPL_DIR = paths.TPL_SG
INBOX = paths.SG_INBOX          # 发起收件箱（与投料口 输入 分开）
# 模板关键词、表头单元格、明细表列布局（1 基）
TPL = {
    "货物": {"kw": "采购", "部门": "C2", "日期": "F2", "预计": "F3", "说明": "C4",
             "cols": (1, 2, 3, 4, 5, 6)},
    "服务": {"kw": "服务", "部门": "C2", "日期": "G2", "预计": "G3", "说明": "C4",
             "cols": (1, 2, 3, 5, 6, 7)},
}


def _fmt_amt(x):
    return str(int(x)) if float(x) == int(float(x)) else str(x)


def _today():
    d = datetime.date.today()
    return f"{d.year}.{d.month}.{d.day}"


def find_approval_row(ws):
    for r in range(ws.max_row, 5, -1):
        v = ws.cell(row=r, column=1).value
        if v and "审批" in str(v):
            return r
    raise RuntimeError("未找到审批意见行，模板结构异常")


def gen(req):
    cat = req.get("类别", "货物")
    if cat not in TPL:
        raise SystemExit(f"类别须为 货物/服务，收到 {cat}")
    name = req["事项名"]
    amt = req.get("金额")
    rows = req.get("明细") or []
    if not name or not amt:
        raise SystemExit("事项名/金额 必填（金额写入文件名，供事项审批取数）")
    if not rows:
        raise SystemExit("明细不能为空")

    cfg = TPL[cat]
    cands = [p for p in TPL_DIR.glob("申购单*.xls*") if cfg["kw"] in p.name]
    if not cands:
        raise SystemExit(f"模板未找到：{TPL_DIR} 中无含“{cfg['kw']}”的申购单模板")
    tpl_file = cands[0]

    wb = load_workbook(tpl_file)
    ws = wb.active
    ap_row = find_approval_row(ws)
    first, last = 6, ap_row - 1  # 明细区：第6行 至 审批块上一行
    if len(rows) > last - first + 1:
        raise SystemExit(f"明细 {len(rows)} 行超出模板容量 {last - first + 1} 行")

    # 表头：申购日期（台账给了就用，否则今天）；申购部门、预计使用日期、申购说明按需覆盖
    ws[cfg["日期"]] = str(req.get("申购日期") or _today())
    if req.get("申购部门"):
        ws[cfg["部门"]] = str(req["申购部门"])
    if req.get("预计使用日期"):
        ws[cfg["预计"]] = str(req["预计使用日期"])
    ws[cfg["说明"]] = req.get("申购说明", name)

    # 清空旧明细并写入新明细
    c_no, c_name, c_spec, c_unit, c_qty, c_note = cfg["cols"]
    for r in range(first, last + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if not isinstance(cell, MergedCell):
                cell.value = None
    for i, item in enumerate(rows):
        r = first + i
        ws.cell(row=r, column=c_no, value=i + 1)
        ws.cell(row=r, column=c_name, value=item.get("名称", ""))
        ws.cell(row=r, column=c_spec, value=item.get("规格", ""))
        ws.cell(row=r, column=c_unit, value=item.get("单位", ""))
        ws.cell(row=r, column=c_qty, value=item.get("数量", ""))
        ws.cell(row=r, column=c_note, value=item.get("备注", ""))
        # 名称/规格过长会折行：按折行数抬高行高，避免固定行高把第二行裁掉
        import math
        disp = lambda s: sum(1 if ord(ch) > 0x2E80 else 0.55 for ch in str(s))
        lines = max(math.ceil(disp(item.get("名称", "")) / 8),
                    math.ceil(disp(item.get("规格", "")) / 9), 1)
        if lines > 1:
            ws.row_dimensions[r].height = 15 * lines + 4

    INBOX.mkdir(parents=True, exist_ok=True)
    out = INBOX / f"申购单（{name}-金额：{_fmt_amt(amt)}）.xlsx"
    wb.save(out)
    print(f"[ok] 生成 {out.name}（{cat}类，{len(rows)}行明细）")
    try:                                    # 交付区转PDF失败不影响申购单主产物
        from oa_common import deliver
        deliver.stage("申购单", [out], name)
    except Exception as e:
        print(f"  [交付] 申购单转PDF失败(不影响生成): {str(e)[:60]}")
    return out


SG_LEDGER = paths.LEDGER_SG


def _colmap(ws):
    """表头名(去掉尾部*) -> 列号。台账挪列/加列/改名(除关键字)都不用动代码。"""
    m = {}
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or '').strip().rstrip('*')
        if h:
            m.setdefault(h, c)
    return m


def read_ledger_groups():
    """读申购单需求台账 -> [(事项名, 组信息)]；状态为空且非示例的行才处理。"""
    wb = load_workbook(SG_LEDGER, data_only=True)
    ws = wb["需求台账"]
    m = _colmap(ws)

    def val(r, key):
        c = m.get(key)
        v = ws.cell(row=r, column=c).value if c else None
        return str(v).strip() if v is not None else ""

    groups, order = {}, []
    for r in range(2, ws.max_row + 1):
        name = val(r, "事项名")
        if not name or val(r, "状态") in ("示例", "已生成"):
            continue
        g = groups.setdefault(name, {"rows": [], "类别": "", "金额": "", "申购说明": "",
                                     "预计使用日期": "", "申购部门": "", "申购日期": ""})
        g["rows"].append({"row": r, "名称": val(r, "物资名称"),
                          "规格": val(r, "规格型号及技术要求"), "单位": val(r, "单位"),
                          "数量": val(r, "数量"), "备注": val(r, "备注")})
        if len(g["rows"]) == 1:
            order.append(name)
            cat = val(r, "类别")
            g["类别"] = "服务" if "服务" in cat else "货物"
            g["金额"] = val(r, "预估金额（元）") or val(r, "预估金额")
            g["申购说明"] = val(r, "申购说明")
            g["预计使用日期"] = val(r, "预计使用日期")
            g["申购部门"] = val(r, "申购部门")
            g["申购日期"] = val(r, "申购日期")
    wb.close()
    return [(name, groups[name]) for name in order]


def write_back(rows, status="已生成"):
    wb = load_workbook(SG_LEDGER)
    ws = wb["需求台账"]
    m = _colmap(ws)
    c_st, c_tm = m.get("状态"), m.get("生成时间")
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    for r in rows:
        if c_st:
            ws.cell(row=r, column=c_st, value=status)
        if c_tm:
            ws.cell(row=r, column=c_tm, value=now)
    wb.save(SG_LEDGER)


def run_ledger():
    groups = read_ledger_groups()
    if not groups:
        print("[ledger] 没有待生成的申购单（状态为空的行）")
        return 0
    ok, fail = 0, 0
    for name, g in groups:
        try:
            req = {"类别": g["类别"] or "货物", "事项名": name, "金额": g["金额"],
                   "申购说明": g["申购说明"] or name,
                   "预计使用日期": g["预计使用日期"],
                   "申购部门": g["申购部门"], "申购日期": g["申购日期"],
                   "明细": [{k: it[k] for k in ("名称", "规格", "单位", "数量", "备注")}
                            for it in g["rows"]]}
            gen(req)
            write_back([it["row"] for it in g["rows"]])
            ok += 1
        except SystemExit as e:
            print(f"[ledger] 组『{name}』生成失败：{e}")
            fail += 1
    print(f"[ledger] 完成：成功{ok}单，失败{fail}单")
    return 0 if fail == 0 else 2


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", help="单条需求 JSON 文件路径")
    ap.add_argument("--ledger", action="store_true",
                    help="批量模式：读取申购单需求台账生成所有待处理事项")
    args = ap.parse_args()
    if args.ledger:
        raise SystemExit(run_ledger())
    if not args.json:
        ap.error("需指定 --json 或 --ledger")
    req = json.loads(Path(args.json).read_text(encoding="utf-8"))
    gen(req)


if __name__ == "__main__":
    main()
