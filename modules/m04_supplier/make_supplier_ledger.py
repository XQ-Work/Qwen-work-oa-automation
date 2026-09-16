# -*- coding: utf-8 -*-
"""生成《供应商建档台账.xlsx》：用户手填/智能体从执照提取 → m04 读取填 OA"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = r"D:\自动备份\Qwen work\自动化（OA流程）\供应商建档台账.xlsx"

HEADERS = [
    ("序号", 6), ("供应商名称*", 26), ("供应商分类*", 14), ("供应商所属类型*", 14),
    ("法定代表人*", 12), ("注册资本*", 14), ("经营范围*", 40),
    ("电话*", 14), ("联系人*", 10), ("开户行*", 24), ("开户行地址", 24),
    ("银行账号*", 22), ("银行行号", 10), ("确定方式", 12), ("情况说明", 24),
    ("状态", 12), ("OA编号", 14), ("建档时间", 16), ("备注", 14),
]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
GRAY = PatternFill("solid", fgColor="D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "建档台账"
for col, (name, width) in enumerate(HEADERS, 1):
    c = ws.cell(row=1, column=col, value=name)
    c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws.column_dimensions[get_column_letter(col)].width = width
ws.row_dimensions[1].height = 30

fills = {2: YELLOW, 3: YELLOW, 4: YELLOW, 5: YELLOW, 6: YELLOW, 7: YELLOW,
         8: YELLOW, 9: YELLOW, 10: YELLOW, 12: YELLOW,
         11: GREEN, 13: GREEN, 14: GREEN, 15: GREEN, 19: GREEN,
         16: GRAY, 17: GRAY, 18: GRAY}

example = ["示例", "上海品凡动力设备有限公司", "货物类", "私企",
           "鲍龙飞", "50万元",
           "动力设备，空气压缩机及配件，空气净化设备，气动设备，机电设备，电动工具，五金交电，建材，管道阀门，电子产品销售，机电设备（除特种设备）安装、维修",
           "021-67606878", "黄勇", "农行上海市航头支行", "农行上海市航头支行",
           "03483900040014873", "\\", "比价", "空压机采购供应商",
           "示例", "", "", "取自历史过审实例"]
for col, val in enumerate(example, 1):
    c = ws.cell(row=2, column=col, value=val)
    c.font = Font(name="微软雅黑", size=10, color="808080", italic=True)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    c.border = BORDER

for r in range(3, 203):
    for col in range(1, len(HEADERS) + 1):
        c = ws.cell(row=r, column=col)
        c.border = BORDER
        c.font = Font(name="微软雅黑", size=10)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if col in fills:
            c.fill = fills[col]
    ws.cell(row=r, column=1).value = f'=IF(B{r}="","",ROW()-1)'

dvs = [("C", '"货物类,服务类,工程类,劳务咨询顾问类"', "供应商分类四选一"),
       ("D", '"国企/央企,私企,政府机关,个人"', "所属类型四选一"),
       ("N", '"公开招标,邀标,比价,单一来源"', "确定方式四选一")]
for letter, formula, err in dvs:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=True, error=err)
    ws.add_data_validation(dv)
    dv.add(f"{letter}3:{letter}202")

ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:S202"

gd = wb.create_sheet("填写说明")
gd.column_dimensions["A"].width = 105
lines = [
    ("《供应商建档台账》填写说明", True),
    ("", False),
    ("■ 一行 = 一家新供应商。黄色必填，绿色选填有默认值，灰色脚本回写别动。", True),
    ("", False),
    ("数据来源分工（也可以全交给智能体）：", True),
    ("  营业执照能提取的：名称/分类判断/法定代表人/注册资本/经营范围/所属类型", False),
    ("  报价单或开票资料能提取的：电话/联系人/开户行/银行账号", False),
    ("  需要你确认补的：开户行地址(默认=开户行同名)、银行行号(可填\\)、确定方式(默认比价)", False),
    ("  把执照+报价单丢进发起收件箱说「建档供应商」，智能体提取后回填本表，你只需核对补空", False),
    ("", False),
    ("附件 6 件套（人工备齐，命名「材料名（供应商全称）.pdf」放发起收件箱）：", True),
    ("  营业执照 / 法人身份证 / 信用中国截图 / 企业公示系统截图 / 政府采购网截图 / 事项审批打印件", False),
    ("  脚本会做齐套检查，缺件提醒不静默", False),
    ("", False),
    ("使用：填好行后说「建档供应商 第N行」或「建档所有空状态行」", True),
    ("  → 脚本读表入库(supplier_profile 主数据) → 登录OA填表 → 传附件 → 存草稿 → 回写状态", False),
    ("  提交仍由人工在 OA 草稿箱完成（公约5）", False),
    ("", False),
    ("说明：公司类型固定选「供应商(我方为甲方)」、操作状态固定「新增」、板块固定「水务」、是否内部公司「否」，脚本内置无需填", False),
]
for i, (text, bold) in enumerate(lines, 1):
    c = gd.cell(row=i, column=1, value=text)
    c.font = Font(name="微软雅黑", size=11 if bold else 10, bold=bold)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print("saved:", OUT)
