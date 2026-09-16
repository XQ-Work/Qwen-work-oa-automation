# -*- coding: utf-8 -*-
"""m04 供应商新增一条龙：读建档台账 -> 入主数据库 -> 登录OA填表 -> 6件套附件齐套检查
-> 存草稿 -> 回写台账。提交永远人工（公约5）。仅支持"新增"操作（用户定）。

用法（仓库根目录）：
  python modules/m04_supplier/fill_supplier.py --row 3            # 干跑：只填不存
  python modules/m04_supplier/fill_supplier.py --row 3 --save     # 正式：填+存草稿+回写
  python modules/m04_supplier/fill_supplier.py --all --save       # 所有空状态行
"""
import argparse
import datetime
import json
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from oa_common import paths, auth, db  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

LEDGER = paths.DATA / "供应商建档台账.xlsx"
WORKFLOWID = "444401"
ADD_PATH = ("/workflow/request/AddRequest.jsp?workflowid=" + WORKFLOWID +
            "&isagent=0&beagenter=0&f_weaver_belongto_userid=")
KEEP_OPEN_SECONDS = 15 * 60

# 建档台账列 -> 键
COLS = {2: "名称", 3: "分类", 4: "所属类型", 5: "法人", 6: "注册资本", 7: "经营范围",
        8: "电话", 9: "联系人", 10: "开户行", 11: "开户行地址", 12: "账号",
        13: "行号", 14: "确定方式", 15: "情况说明", 16: "状态"}

# OA 字段编号（见 FIELD_MAP_供应商.md）
F = {"公司类型": "field403694", "操作状态": "field403696", "板块": "field403640",
     "确定方式": "field507131", "情况说明": "field511130", "名称": "field403642",
     "所属类型": "field403645", "法人": "field403643", "注册资本": "field403644",
     "经营范围": "field403650", "是否内部公司": "field403646", "电话": "field403664",
     "联系人": "field403674", "开户行": "field403702_0", "开户行地址": "field403705_0",
     "账号": "field403708_0", "行号": "field403711_0"}
CATEGORY_CHECKBOX = {"货物类": "field403636", "服务类": "field403637",
                     "工程类": "field403638", "劳务咨询顾问类": "field507130"}
# 6 件套：材料关键词
ANNEX_KINDS = ["营业执照", "身份证", "信用中国", "公示", "政府采购", "事项审批"]


def norm(s):
    return re.sub(r"[（）()【】\[\]\s\-_·、,，/]", "", s or "")


def read_row(r):
    wb = load_workbook(LEDGER, data_only=True)
    ws = wb["建档台账"]
    rec = {k: str(ws.cell(row=r, column=c).value or "").strip()
           for c, k in COLS.items()}
    wb.close()
    if not rec["名称"]:
        raise SystemExit(f"台账第{r}行供应商名称为空")
    return rec


def pending_rows():
    wb = load_workbook(LEDGER, data_only=True)
    ws = wb["建档台账"]
    out = []
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(row=r, column=2).value or "").strip()
        st = str(ws.cell(row=r, column=16).value or "").strip()
        if name and st not in ("示例", "已存草稿", "已提交"):
            out.append(r)
    wb.close()
    return out


def find_form_frame(page):
    """CDP content() 普查：field* 控件最多的帧（该表单帧 JS 环境不可靠）"""
    for _ in range(15):
        best = None
        for f in page.frames:
            try:
                html = f.content()
            except Exception:
                continue
            n = html.count('name="field') if html else 0
            if n >= 5 and (best is None or n > best[0]):
                best = (n, f)
        if best:
            return best[1]
        time.sleep(2)
    return None


def match_annex_files(supplier):
    """在收件箱找该供应商的附件，返回 {材料类型: 路径} 与缺失列表"""
    ns = norm(supplier)
    found, missing = {}, []
    files = [p for p in paths.INBOX.glob("*") if p.suffix.lower()
             in (".pdf", ".png", ".jpg", ".jpeg", ".doc", ".docx")]
    for kind in ANNEX_KINDS:
        hit = None
        for p in files:
            if kind not in p.stem:
                continue
            parens = re.findall(r"[（(]([^（）()]+)[）)]", p.stem)
            no = norm("".join(parens))
            owner_ok = bool(no) and (no in ns or ns in no or ns.startswith(no))
            # 事项审批打印件历史上按事项命名（括号内非供应商名），只认关键词
            if owner_ok or (kind == "事项审批" and not no):
                hit = str(p)
                break
            if kind == "事项审批" and no and no not in ns and ns not in no:
                # 括号是事项名：也接受（建档场景收件箱即本批次材料）
                hit = str(p)
                break
        if hit:
            found[kind] = hit
        else:
            missing.append(kind)
    return found, missing


def build_plan(rec):
    """填值计划：(kind, 字段, 期望值, 描述)。父级下拉在前，子级字段在后（对抗级联刷新）"""
    situation = rec["情况说明"] or f"{rec['名称'][:6]}供应商"
    cb = CATEGORY_CHECKBOX.get(rec["分类"])
    if not cb:
        raise SystemExit(f"分类『{rec['分类']}』不识别（货物类/服务类/工程类/劳务咨询顾问类）")
    return [
        ("select", F["公司类型"], "供应商（我方为甲方）", "公司类型"),
        ("select", F["操作状态"], "新增", "操作状态"),
        ("select", F["板块"], "水务板块", "所属板块"),
        ("select", F["是否内部公司"], "否", "是否内部公司"),
        ("select", F["确定方式"], rec["确定方式"] or "比价", "确定方式"),
        ("select", F["所属类型"], rec["所属类型"] or "私企", "所属类型"),
        ("check", cb, "1", f"分类={rec['分类']}"),
        ("text", F["名称"], rec["名称"], "供应商名称"),
        ("text", F["法人"], rec["法人"], "法定代表人"),
        ("text", F["注册资本"], rec["注册资本"], "注册资本"),
        ("area", F["经营范围"], rec["经营范围"], "经营范围"),
        ("text", F["电话"], rec["电话"], "电话"),
        ("text", F["联系人"], rec["联系人"], "联系人"),
        ("area", F["情况说明"], situation, "情况说明"),
        ("text", F["开户行"], rec["开户行"], "开户行"),
        ("text", F["开户行地址"], rec["开户行地址"] or rec["开户行"], "开户行地址"),
        ("text", F["账号"], rec["账号"], "银行账号"),
        ("text", F["行号"], rec["行号"] or "\\", "行号"),
    ]


def apply_item(frame, kind, name, expected):
    if kind == "select":
        frame.locator(f"select[name={name}]:visible").first.select_option(label=expected)
    elif kind == "area":
        frame.locator(f"textarea[name={name}]").first.fill(expected)
    elif kind == "text":
        frame.locator(f"input[name={name}]:visible").first.fill(expected)
    elif kind == "check":
        loc = frame.locator(f"input[name={name}]").first
        try:  # 先点 jNice 样式元素（视觉同步）
            if not loc.is_checked():
                sib = frame.locator(f"input[name={name}] + *").first
                if sib.count() > 0:
                    sib.click(timeout=3000)
        except Exception:
            pass
        try:
            if not loc.is_checked():
                loc.evaluate("el => { el.checked = true; "
                             "el.dispatchEvent(new Event('change',{bubbles:true})); }")
        except Exception:
            pass


def verify_item(frame, kind, name, expected):
    try:
        if kind == "check":
            return frame.locator(f"input[name={name}]").first.is_checked()
        if kind == "select":
            v = frame.locator(f"select[name={name}]:visible").first.input_value()
            if v == expected:
                return True
            try:  # input_value 是选项value，比对选中项文本
                t = frame.locator(
                    f"select[name={name}]:visible option:checked").first.inner_text()
                return t.strip() == expected.strip()
            except Exception:
                return False
        v = frame.locator(
            f"input[name={name}]:visible, textarea[name={name}]").first.input_value()
        return v.strip() == expected.strip()
    except Exception:
        return False


def fill_form(frame, rec, step):
    plan = build_plan(rec)
    for kind, name, expected, desc in plan:
        apply_item(frame, kind, name, expected)
    step("首轮填报", True, f"{len(plan)}项")
    time.sleep(4)  # 等级联刷新收敛
    # 多轮读回补填（级联会清空子级字段）
    bad = [it for it in plan if not verify_item(frame, it[0], it[1], it[2])]
    for rnd in range(3):
        if not bad:
            break
        time.sleep(2)
        for kind, name, expected, desc in bad:
            apply_item(frame, kind, name, expected)
        time.sleep(3)
        bad = [it for it in plan if not verify_item(frame, it[0], it[1], it[2])]
    step("读回校验补填", not bad,
         f"全部{len(plan)}项就位" if not bad
         else "仍未就位: " + "、".join(it[3] for it in bad))


def to_db(rec):
    conn = db.connect()
    db.upsert_profile(conn, rec["名称"], {
        "category": rec["分类"], "type": rec["所属类型"], "legal_person": rec["法人"],
        "reg_capital": rec["注册资本"], "scope": rec["经营范围"],
        "phone": rec["电话"], "contact": rec["联系人"], "confirm_method": rec["确定方式"],
        "situation": rec["情况说明"], "bank_name": rec["开户行"],
        "bank_addr": rec["开户行地址"], "bank_account": rec["账号"],
        "bank_code": rec["行号"]}, source="建档台账")
    conn.commit()


def write_back(r, status, reqid=""):
    wb = load_workbook(LEDGER)
    ws = wb["建档台账"]
    ws.cell(row=r, column=16, value=status)
    if reqid:
        ws.cell(row=r, column=17, value=reqid)
    ws.cell(row=r, column=18,
            value=datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))
    wb.save(LEDGER)


def process(r, save):
    rec = read_row(r)
    report = {"row": r, "rec": rec, "steps": []}

    def step(name, ok, err=""):
        report["steps"].append({"step": name, "ok": bool(ok), "err": str(err)[:150]})
        print(f"[{'ok ' if ok else 'ERR'}] {name} {err}", flush=True)

    to_db(rec)
    step("画像入主数据库", True)
    annex_files, missing = match_annex_files(rec["名称"])
    step("附件齐套检查", not missing,
         f"有{len(annex_files)}件" + (f"，缺：{missing}" if missing else "，齐全"))

    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(paths.PROFILE), channel="msedge", headless=False,
            args=["--start-maximized"], no_viewport=True)
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        try:
            base = auth.login(page)
            step("登录", True, base)
            page.goto(base + ADD_PATH, wait_until="domcontentloaded", timeout=40000)
            frame = find_form_frame(page)
            if frame is None:
                step("定位表单帧", False)
                raise RuntimeError("未找到表单帧")
            step("定位表单帧", True, frame.url[:70])
            fill_form(frame, rec, step)
            if annex_files:
                btn = frame.get_by_text("选取多个文件").first
                try:
                    with page.expect_file_chooser(timeout=15000) as fc:
                        btn.click()
                    fc.value.set_files(list(annex_files.values()))
                    step("上传附件", True, f"{len(annex_files)}个")
                    time.sleep(5)
                except Exception as e:
                    step("上传附件", False, e)
            page.screenshot(path=str(paths.STATE / "supplier_filled.png"),
                            full_page=True)
            step("整页截图", True, "supplier_filled.png")

            if save:
                signals = []
                page.on("dialog", lambda d: (signals.append(d.message or ""), d.accept()))
                try:
                    page.locator("input.e8_btn_top[value=保存]").first.click(timeout=8000)
                    step("点保存", True)
                except Exception as e:
                    step("点保存", False, e)
                reqid, ok = "", False
                for _ in range(20):
                    time.sleep(1.5)
                    m = re.search(r"requestid=(\d+)", page.url or "", re.I)
                    if m:
                        reqid, ok = m.group(1), True
                        break
                    if any("成功" in s for s in signals):
                        ok = True
                        break
                page.screenshot(path=str(paths.STATE / "supplier_saved.png"),
                                full_page=True)
                step("保存草稿", ok, f"requestid={reqid}" if reqid else "未见编号,看supplier_saved.png")
                if ok:
                    write_back(r, "已存草稿", reqid)
                    conn = db.connect()
                    conn.execute(
                        "INSERT INTO oa_flow(matter_id,flow_type,requestid,title,status,created_at)"
                        " VALUES(NULL,?,?,?,?,?)",
                        ("供应商新增", reqid, rec["名称"], "已存草稿", db.now()))
                    conn.commit()
                    # 附件移入供应商资料档
                    tgt = paths.INBOX.parent / "供应商资料" / re.sub(r"[\\/:*?\"<>|]", "_", rec["名称"])
                    tgt.mkdir(parents=True, exist_ok=True)
                    import shutil
                    for f in annex_files.values():
                        shutil.move(f, str(tgt / Path(f).name))
                    step("附件归档供应商资料", True, str(tgt))
                return ok
            else:
                print(f"[hold] 干跑：浏览器保持 {KEEP_OPEN_SECONDS//60} 分钟供核对，不保存。",
                      flush=True)
                time.sleep(KEEP_OPEN_SECONDS)
                return True
        finally:
            (paths.STATE / f"supplier_row{r}_report.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            try:
                ctx.close()
            except Exception:
                pass


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--row", type=int)
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--save", action="store_true")
    args = ap.parse_args()
    rows = [args.row] if args.row else pending_rows()
    if not rows:
        print("没有待建档行")
        return
    ok_rows = []
    for r in rows:
        print(f"########## 建档台账 行{r} ##########", flush=True)
        try:
            if process(r, args.save):
                ok_rows.append(r)
        except SystemExit as e:
            print(f"[skip] 行{r}: {e}", flush=True)
        except Exception as e:
            print(f"[ERR] 行{r}: {e}", flush=True)
        time.sleep(3)
    print(f"完成：成功{len(ok_rows)}/{len(rows)}")


if __name__ == "__main__":
    main()
