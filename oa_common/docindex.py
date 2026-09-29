# -*- coding: utf-8 -*-
"""文档索引（投料即建档 · 第一阶段：文字版抽取，不做 OCR）。

intake run 归档每个文件时调用 ingest()：抽正文入 FTS5 全文索引，
同 sha 秒跳过；图片/扫描件只登记元数据标 status=待视觉（第二阶段接 OCR）。

表（独立库 记录/db/文档索引.db，正文含敏感信息，与台账同级永不入 git）：
  docs       文件档案行：sha/路径/材料名/项目编号/提取方式/状态/关键字段meta
  docs_fts   FTS5 全文：title + body
  ocr_cache  按图片内容MD5缓存识别结果，同页不重跑

CLI：
  python -m oa_common.docindex stat                     索引概况
  python -m oa_common.docindex search 关键词 [编号]      全文检索（可限项目）
  python -m oa_common.docindex add <文件> [材料名]       手动补录单个文件
  python -m oa_common.docindex rescan                    增量补录 记录/项目档案 全部文字版
  python -m oa_common.docindex ocr [--limit N]           二期：对 待视觉/文字稀少 件跑OCR
"""
import datetime as dt
import hashlib
import json
import re
import sqlite3
import sys
import zipfile
from pathlib import Path

from oa_common import paths

DB = paths.DOCS_DB
STAGE_MATERIAL = {"01": "事项审批", "02": "报价单及资质", "03": "比质比价报告",
                  "04": "合同", "05": "验收资料", "06": "付款单"}
MIN_TEXT = 150          # 少于该字符数视为疑似扫描件
BODY_CAP = 120000       # 正文入库上限

_SCHEMA = """
CREATE TABLE IF NOT EXISTS docs(
  sha TEXT PRIMARY KEY, path TEXT, material TEXT, code TEXT, matter TEXT,
  kind TEXT, status TEXT, chars INTEGER DEFAULT 0, meta TEXT DEFAULT '',
  mtime TEXT, created_at TEXT);
CREATE VIRTUAL TABLE IF NOT EXISTS docs_fts USING fts5(title, body, sha UNINDEXED);
CREATE TABLE IF NOT EXISTS ocr_cache(
  key TEXT PRIMARY KEY, text TEXT, created_at TEXT);
"""

# 材料名 -> 正文必含特征（防错投：文件名与内容不符报警）
SNIFF = {
    "发票": ["发票", "税额"],
    "合同": ["合同", "甲方"],
    "付款单": ["付款", "金额"],
    "验收资料": ["验收"],
    "申购单": ["申购"],
    "事项审批": ["审批", "事由"],
    "比质比价报告": ["比价", "报价"],
    "报价单及资质": ["报价", "价"],
}


def connect():
    DB.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(DB))
    conn.row_factory = sqlite3.Row
    conn.executescript(_SCHEMA)
    return conn


def now():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------- 各类型抽取（第一阶段：只碰文字层） ----------

def _pdf_text(path):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    with fitz.open(str(path)) as doc:
        return "\n".join(pg.get_text() for pg in doc)


def _docx_text(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
    xml = re.sub(r"<w:p[ >]", "\n<w:p ", xml)
    return re.sub(r"<[^>]+>", "", xml)


def _xlsx_text(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 200:
                break
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def extract(path):
    """返回 (kind, text, meta)；图片返回 ('image', '', {})。"""
    p = Path(path)
    suf = p.suffix.lower()
    if suf == ".pdf":
        try:
            text = _pdf_text(p)
        except Exception:
            return "pdf_err", "", {}
        if re.match(r"^dzfp_|数电票", p.name) and "发票号码" in text:
            try:
                sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "modules/m06_payment"))
                import parse_invoice
                inv = parse_invoice.parse(p)
                return "invoice", text, {k: str(v) for k, v in inv.items() if v}
            except Exception:
                pass
        return "pdf_text", text, {}
    if suf == ".docx":
        return "docx_text", _docx_text(p), {}
    if suf == ".xlsx":
        return "xlsx_text", _xlsx_text(p), {}
    if suf in (".png", ".jpg", ".jpeg"):
        return "image", "", {}
    if suf in (".doc", ".xls", ".lnk"):
        return "legacy", "", {}
    return "other", "", {}


def _norm(s):
    return re.sub(r"\s+", "", s or "")


def sniff(material, text):
    """错投检测：按材料名核对正文特征，返回 None(正常) 或 提示语。"""
    kws = SNIFF.get(material)
    if not kws or len(_norm(text)) < MIN_TEXT:
        return None
    body = _norm(text)
    if not any(k in body for k in kws):
        return f"正文未出现「{material}」特征词({'/'.join(kws)})，疑似错投"
    return None


def ingest(path, material="", code="", matter=""):
    """登记一个已归档文件。返回 dict(sha/status/warn)。幂等：同 sha 直接跳过。"""
    p = Path(path)
    if not p.exists():
        return {"sha": "", "status": "missing", "warn": "文件不存在"}
    sha = sha256(p)
    conn = connect()
    if conn.execute("SELECT 1 FROM docs WHERE sha=?", (sha,)).fetchone():
        conn.close()
        return {"sha": sha, "status": "skip", "warn": ""}
    kind, text, meta = extract(p)
    text = text[:BODY_CAP]
    status = "已登记"
    warn = ""
    if len(_norm(text)) < MIN_TEXT:
        status = "待视觉" if kind in ("image", "pdf_err", "legacy", "other") else "文字稀少"
    else:
        warn = sniff(material, text) or ""
        if warn:
            status = "名实存疑"
    mtime = dt.datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
    conn.execute("INSERT OR REPLACE INTO docs VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                 (sha, str(p.resolve()), material, code, matter, kind, status,
                  len(text), json.dumps(meta, ensure_ascii=False), mtime, now()))
    conn.execute("DELETE FROM docs_fts WHERE sha=?", (sha,))
    conn.execute("INSERT INTO docs_fts(title, body, sha) VALUES(?,?,?)",
                 (p.name, text, sha))
    conn.commit()
    conn.close()
    return {"sha": sha, "status": status, "warn": warn}


def search(kw, code=""):
    conn = connect()
    sql = ("SELECT d.path, d.material, d.code, d.status, "
           "snippet(docs_fts, 1, '[', ']', '…', 20) AS snip "
           "FROM docs_fts JOIN docs d ON d.sha=docs_fts.sha "
           "WHERE docs_fts MATCH ? ")
    args = [kw]
    if code:
        sql += "AND d.code=? "
        args.append(code)
    sql += "ORDER BY rank LIMIT 30"
    try:
        rows = conn.execute(sql, args).fetchall()
    except sqlite3.OperationalError:      # 中文分词简单兜底
        rows = conn.execute(
            "SELECT path, material, code, status, substr('',1,0) FROM docs "
            "WHERE sha IN (SELECT sha FROM docs_fts WHERE body LIKE ?) LIMIT 30",
            ("%" + kw + "%",)).fetchall()
    conn.close()
    for r in rows:
        snip = (r["snip"] or "")[:80]
        print(f"[{r['material'] or '?'}] {r['code'] or ''} {Path(r['path']).name}  ({r['status']})")
        if snip:
            print(f"    …{snip}…")
    print(f"命中 {len(rows)} 条")
    return rows


def stat():
    conn = connect()
    rows = conn.execute("SELECT status, COUNT(*) n, SUM(chars) c FROM docs GROUP BY status").fetchall()
    print("索引库:", DB)
    for r in rows:
        print(f"  {r['status']}: {r['n']} 件（正文共 {r['c'] or 0} 字）")
    if not rows:
        print("  （空）")
    conn.close()


def rescan():
    """增量补录 记录/项目档案 下全部文字版文件（跳过已登记 sha）。"""
    n = done = 0
    if paths.ARCHIVE.is_dir():
        for f in paths.ARCHIVE.rglob("*"):
            if not f.is_file() or f.suffix.lower() not in (
                    ".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg"):
                continue
            code, matter = "", ""
            m = re.match(r"^(\d{4}-\d{3})-(.+)$", f.parent.name)
            material = STAGE_MATERIAL.get(f.parent.name[:2], "")
            if m:
                code, matter = m.group(1), m.group(2)
            r = ingest(f, material=material, code=code, matter=matter)
            n += 1
            if r["status"] != "skip":
                done += 1
    print(f"扫描 {n} 个文件，新登记 {done}，其余已存在跳过")


# ---------- 二期：OCR（Windows 内置引擎，复用单据类 ocr_local.ps1） ----------

OCR_PS = Path(__file__).resolve().parents[1] / "danju" / "80_工具" / "ocr_local.ps1"
OCR_DPI = 170
OCR_MAX_PAGES = 30


def _img_md5(p):
    h = hashlib.md5()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def _ocr_image(img):
    """单张图 OCR，带 ocr_cache（按内容MD5）。返回文本或 None。"""
    conn = connect()
    key = _img_md5(img)
    row = conn.execute("SELECT text FROM ocr_cache WHERE key=?", (key,)).fetchone()
    conn.close()
    if row is not None:
        return row["text"]
    if not OCR_PS.exists():
        return None
    import subprocess
    try:
        r = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
             "-File", str(OCR_PS), str(img)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=120)
    except Exception:
        return None
    out = r.stdout or ""
    body = re.sub(r"^-{40,}$", "", out, flags=re.M).strip()
    if "[ERROR]" in body or not body:
        return None
    conn = connect()
    conn.execute("INSERT OR REPLACE INTO ocr_cache VALUES(?,?,?)", (key, body, now()))
    conn.commit()
    conn.close()
    return body


def _ocr_pdf(pdf):
    """无文字层PDF：逐页渲染后OCR（页数封顶）。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    tmp = paths.STATE / "ocr_tmp_docindex"
    tmp.mkdir(parents=True, exist_ok=True)
    texts = []
    with fitz.open(str(pdf)) as doc:
        for i, pg in enumerate(doc):
            if i >= OCR_MAX_PAGES:
                break
            out = tmp / f"{pdf.stem}_p{i+1}".replace(" ", "_")
            out = out.with_suffix(".png")
            pg.get_pixmap(dpi=OCR_DPI).save(out)
            t = _ocr_image(out)
            try:
                out.unlink()
            except OSError:
                pass
            if t:
                texts.append(t)
    return "\n".join(texts)


def _set_body(sha, text, kind, status):
    conn = connect()
    conn.execute("UPDATE docs SET kind=?, status=?, chars=? WHERE sha=?",
                 (kind, status, len(text), sha))
    conn.execute("DELETE FROM docs_fts WHERE sha=?", (sha,))
    path = conn.execute("SELECT path FROM docs WHERE sha=?", (sha,)).fetchone()["path"]
    conn.execute("INSERT INTO docs_fts(title, body, sha) VALUES(?,?,?)",
                 (Path(path).name, text[:BODY_CAP], sha))
    conn.commit()
    conn.close()


def ocr_backfill(limit=None):
    """对待视觉/文字稀少 的档案行跑OCR补正文。逐行可断点续跑（成功即改状态）。"""
    conn = connect()
    rows = conn.execute(
        "SELECT sha, path FROM docs WHERE status IN ('待视觉','文字稀少')").fetchall()
    conn.close()
    prio = {".png": 0, ".jpg": 0, ".jpeg": 0, ".pdf": 1}
    rows = sorted(rows, key=lambda r: prio.get(Path(r["path"]).suffix.lower(), 9))
    if limit:
        rows = rows[:limit]
    ok = fail = skip = 0
    for i, r in enumerate(rows, 1):
        p = Path(r["path"])
        if not p.exists():
            fail += 1
            continue
        suf = p.suffix.lower()
        if suf not in (".png", ".jpg", ".jpeg", ".pdf"):
            conn = connect()     # doc/xls/lnk 等一期不OCR：落终态免重复重试
            conn.execute("UPDATE docs SET status='无需OCR' WHERE sha=?", (r["sha"],))
            conn.commit()
            conn.close()
            skip += 1
            continue
        try:
            text = _ocr_image(p) or "" if suf != ".pdf" else _ocr_pdf(p)
        except Exception:
            text = ""
        if len(_norm(text)) >= MIN_TEXT // 2:
            _set_body(r["sha"], text, "ocr", "已登记(OCR)")
            ok += 1
            print(f"  [{i}/{len(rows)}] OCR成功 {p.name}（{len(text)}字）")
        else:
            fail += 1
            print(f"  [{i}/{len(rows)}] OCR无有效结果 {p.name}")
    print(f"OCR补录完成：成功 {ok}，未成功 {fail}（保持待视觉可重跑），跳过非图格式 {skip}")


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    cmd = sys.argv[1] if len(sys.argv) > 1 else "stat"
    if cmd == "stat":
        stat()
    elif cmd == "search" and len(sys.argv) > 2:
        search(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else "")
    elif cmd == "add" and len(sys.argv) > 2:
        print(ingest(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else ""))
    elif cmd == "rescan":
        rescan()
    elif cmd == "ocr":
        lim = None
        if "--limit" in sys.argv:
            try:
                lim = int(sys.argv[sys.argv.index("--limit") + 1])
            except (IndexError, ValueError):
                pass
        ocr_backfill(limit=lim)
    else:
        print(__doc__)
